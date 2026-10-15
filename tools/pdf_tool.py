# -*- coding: utf-8 -*-
"""
PDF全能工具箱模块
功能：合并PDF、拆分PDF、提取文本、PDF转图片、PDF转Word
"""
import customtkinter as ctk
from tkinter import filedialog, messagebox
import os
from PyPDF2 import PdfReader, PdfWriter
from docx import Document
import fitz  # PyMuPDF


class PDFToolFrame(ctk.CTkFrame):
    """PDF工具箱 - CustomTkinter Frame版本"""

    def __init__(self, master=None, **kwargs):
        super().__init__(master, **kwargs)
        self.merge_files = []
        self.create_widgets()

    def create_widgets(self):
        # 标题
        title = ctk.CTkLabel(
            self,
            text="📄 PDF 全能工具箱",
            font=ctk.CTkFont(size=20, weight="bold")
        )
        title.pack(pady=(5, 15))

        # 标签页
        self.tabview = ctk.CTkTabview(self, corner_radius=10)
        self.tabview.pack(fill="both", expand=True, padx=10, pady=(0, 10))

        # 5个功能标签
        self.tab_merge = self.tabview.add("🔗 合并PDF")
        self.tab_split = self.tabview.add("✂️ 拆分PDF")
        self.tab_text = self.tabview.add("📝 提取文本")
        self.tab_img = self.tabview.add("🖼️ PDF转图片")
        self.tab_word = self.tabview.add("📄 PDF转Word")

        # 初始化各标签页
        self._init_merge_tab()
        self._init_split_tab()
        self._init_text_tab()
        self._init_img_tab()
        self._init_word_tab()

    def log(self, widget, text):
        """输出日志到指定文本框"""
        widget.insert(ctk.END, text + "\n")
        widget.see(ctk.END)

    # ==================== 1. 合并PDF ====================
    def _init_merge_tab(self):
        frame = ctk.CTkFrame(self.tab_merge, fg_color="transparent")
        frame.pack(fill="both", expand=True, padx=15, pady=15)

        ctk.CTkLabel(frame, text="选择多个PDF文件进行合并", font=ctk.CTkFont(size=13, weight="bold")).pack(anchor="w", pady=(0, 8))

        self.merge_listbox = ctk.CTkTextbox(frame, height=120)
        self.merge_listbox.pack(fill="x", pady=5)

        btn_frame = ctk.CTkFrame(frame, fg_color="transparent")
        btn_frame.pack(pady=10)

        ctk.CTkButton(btn_frame, text="➕ 添加PDF", command=self._add_merge_pdf, width=100).pack(side="left", padx=5)
        ctk.CTkButton(btn_frame, text="🗑️ 清空列表", command=self._clear_merge, width=100, fg_color="gray").pack(side="left", padx=5)
        ctk.CTkButton(btn_frame, text="🚀 开始合并", command=self._merge_pdf, width=100).pack(side="left", padx=5)

        self.merge_log = ctk.CTkTextbox(frame, height=150)
        self.merge_log.pack(fill="both", expand=True, pady=5)
        self.log(self.merge_log, "✅ 添加PDF文件，点击合并即可生成新文件")

    def _add_merge_pdf(self):
        files = filedialog.askopenfilenames(filetypes=[("PDF文件", "*.pdf")])
        if files:
            self.merge_files.extend(files)
            for f in files:
                self.merge_listbox.insert(ctk.END, os.path.basename(f) + "\n")

    def _clear_merge(self):
        self.merge_files = []
        self.merge_listbox.delete(1.0, ctk.END)

    def _merge_pdf(self):
        if not self.merge_files:
            messagebox.showerror("错误", "请先添加PDF文件！")
            return
        save_path = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF文件", "*.pdf")],
            initialfile="合并结果.pdf"
        )
        if not save_path:
            return

        self.log(self.merge_log, "🔄 开始合并...")
        try:
            writer = PdfWriter()
            for file in self.merge_files:
                reader = PdfReader(file)
                for page in reader.pages:
                    writer.add_page(page)
            with open(save_path, "wb") as f:
                writer.write(f)
            self.log(self.merge_log, f"🎉 合并完成！保存至：{os.path.basename(save_path)}")
            messagebox.showinfo("成功", "PDF合并完成！")
        except Exception as e:
            self.log(self.merge_log, f"❌ 失败：{str(e)}")
            messagebox.showerror("失败", f"合并出错：{str(e)}")

    # ==================== 2. 拆分PDF ====================
    def _init_split_tab(self):
        frame = ctk.CTkFrame(self.tab_split, fg_color="transparent")
        frame.pack(fill="both", expand=True, padx=15, pady=15)

        # 文件选择
        file_frame = ctk.CTkFrame(frame, fg_color="transparent")
        file_frame.pack(fill="x", pady=5)
        ctk.CTkLabel(file_frame, text="PDF文件：", width=90).pack(side="left")
        self.split_path = ctk.StringVar()
        ctk.CTkEntry(file_frame, textvariable=self.split_path).pack(side="left", fill="x", expand=True, padx=5)
        ctk.CTkButton(file_frame, text="选择", command=self._select_split_pdf, width=80).pack(side="left")

        # 页码输入
        page_frame = ctk.CTkFrame(frame, fg_color="transparent")
        page_frame.pack(fill="x", pady=10)
        ctk.CTkLabel(page_frame, text="拆分页码：", width=90).pack(side="left")
        self.split_page = ctk.CTkEntry(page_frame, placeholder_text="例如：1-5,7,10-15")
        self.split_page.pack(side="left", fill="x", expand=True, padx=5)

        ctk.CTkLabel(frame, text="💡 支持格式：1-5（连续页）、7（单页）、1-5,7,10-15（组合）",
                     font=ctk.CTkFont(size=12), text_color="#888").pack(anchor="w", pady=(0, 10))

        ctk.CTkButton(frame, text="🚀 开始拆分", command=self._split_pdf, height=36).pack(pady=5)

        self.split_log = ctk.CTkTextbox(frame, height=180)
        self.split_log.pack(fill="both", expand=True, pady=10)
        self.log(self.split_log, "✅ 选择PDF → 输入页码 → 点击拆分")

    def _select_split_pdf(self):
        path = filedialog.askopenfilename(filetypes=[("PDF文件", "*.pdf")])
        if path:
            self.split_path.set(path)

    def _split_pdf(self):
        path = self.split_path.get()
        pages = self.split_page.get()
        if not path or not pages:
            messagebox.showerror("错误", "请填写完整信息！")
            return
        save_dir = filedialog.askdirectory(title="选择保存文件夹")
        if not save_dir:
            return

        self.log(self.split_log, "🔄 开始拆分...")
        try:
            reader = PdfReader(path)
            total_pages = len(reader.pages)
            writer = PdfWriter()

            for part in pages.split(","):
                part = part.strip()
                if "-" in part:
                    s, e = map(int, part.split("-"))
                    for i in range(s - 1, e):
                        if i < total_pages:
                            writer.add_page(reader.pages[i])
                else:
                    i = int(part) - 1
                    if i < total_pages:
                        writer.add_page(reader.pages[i])

            save_path = os.path.join(save_dir, "拆分结果.pdf")
            with open(save_path, "wb") as f:
                writer.write(f)
            self.log(self.split_log, f"🎉 拆分完成！保存至：{save_path}")
            messagebox.showinfo("成功", "PDF拆分完成！")
        except Exception as e:
            self.log(self.split_log, f"❌ 失败：{str(e)}")
            messagebox.showerror("失败", f"拆分出错：{str(e)}")

    # ==================== 3. 提取文本 ====================
    def _init_text_tab(self):
        frame = ctk.CTkFrame(self.tab_text, fg_color="transparent")
        frame.pack(fill="both", expand=True, padx=15, pady=15)

        file_frame = ctk.CTkFrame(frame, fg_color="transparent")
        file_frame.pack(fill="x", pady=5)
        self.text_path = ctk.StringVar()
        ctk.CTkEntry(file_frame, textvariable=self.text_path).pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(file_frame, text="选择PDF", command=self._select_text_pdf, width=100).pack(side="left")

        ctk.CTkButton(frame, text="📝 提取文本", command=self._extract_text, height=36).pack(pady=10)

        self.text_content = ctk.CTkTextbox(frame)
        self.text_content.pack(fill="both", expand=True, pady=5)
        self.log(self.text_content, "✅ 选择PDF后一键提取所有文本内容")

    def _select_text_pdf(self):
        path = filedialog.askopenfilename(filetypes=[("PDF文件", "*.pdf")])
        if path:
            self.text_path.set(path)

    def _extract_text(self):
        path = self.text_path.get()
        if not path:
            messagebox.showerror("错误", "请选择PDF文件！")
            return
        self.text_content.delete(1.0, ctk.END)
        try:
            # 优先使用 PyMuPDF（对中文支持更好）
            doc = fitz.open(path)
            text = ""
            for i, page in enumerate(doc):
                page_text = page.get_text()
                if page_text and page_text.strip():
                    text += f"===== 第 {i+1} 页 =====\n{page_text}\n\n"
            if not text.strip():
                # 如果PyMuPDF提取为空，回退到PyPDF2
                reader = PdfReader(path)
                for i, page in enumerate(reader.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text += f"===== 第 {i+1} 页 =====\n{page_text}\n\n"
            doc.close()
            self.text_content.insert(ctk.END, text)
            self.log(self.text_content, f"文本提取完成！")
        except Exception as e:
            self.log(self.text_content, f"失败：{str(e)}")
            messagebox.showerror("失败", f"提取出错：{str(e)}")

    # ==================== 4. PDF转图片 ====================
    def _init_img_tab(self):
        frame = ctk.CTkFrame(self.tab_img, fg_color="transparent")
        frame.pack(fill="both", expand=True, padx=15, pady=15)

        ctk.CTkLabel(frame, text="将PDF每一页转换为高清PNG图片（2倍清晰度）",
                     font=ctk.CTkFont(size=13), text_color="#666").pack(anchor="w", pady=(0, 10))

        file_frame = ctk.CTkFrame(frame, fg_color="transparent")
        file_frame.pack(fill="x", pady=5)
        self.pdf_img_path = ctk.StringVar()
        ctk.CTkEntry(file_frame, textvariable=self.pdf_img_path).pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(file_frame, text="选择PDF", command=self._select_pdf_img, width=100).pack(side="left")

        ctk.CTkButton(frame, text="🖼️ 转换为图片", command=self._pdf_to_image, height=40).pack(pady=15)

        self.img_log = ctk.CTkTextbox(frame, height=200)
        self.img_log.pack(fill="both", expand=True, pady=5)
        self.log(self.img_log, "✅ PDF每一页 → 高清PNG图片")
        self.log(self.img_log, "💡 2倍清晰度渲染，无损保存")

    def _select_pdf_img(self):
        path = filedialog.askopenfilename(filetypes=[("PDF文件", "*.pdf")])
        if path:
            self.pdf_img_path.set(path)

    def _pdf_to_image(self):
        pdf_path = self.pdf_img_path.get()
        if not pdf_path:
            messagebox.showerror("错误", "请选择PDF文件！")
            return

        save_dir = filedialog.askdirectory(title="选择图片保存文件夹")
        if not save_dir:
            return

        self.log(self.img_log, "🔄 正在转换PDF为图片...")
        try:
            doc = fitz.open(pdf_path)
            total = len(doc)
            for page_num in range(total):
                page = doc[page_num]
                # 2倍高清渲染
                mat = fitz.Matrix(2.0, 2.0)
                pix = page.get_pixmap(matrix=mat)
                img_path = os.path.join(save_dir, f"第{page_num+1}页.png")
                pix.save(img_path)
                self.log(self.img_log, f"  ✓ 第 {page_num+1}/{total} 页完成")

            self.log(self.img_log, f"🎉 转换完成！共 {total} 页高清图片")
            messagebox.showinfo("成功", f"PDF转图片完成！\n共{total}页高清图片")
        except Exception as e:
            self.log(self.img_log, f"❌ 转换失败：{str(e)}")
            messagebox.showerror("失败", f"转换出错：{str(e)}")

    # ==================== 5. PDF转Word ====================
    def _init_word_tab(self):
        frame = ctk.CTkFrame(self.tab_word, fg_color="transparent")
        frame.pack(fill="both", expand=True, padx=15, pady=15)

        ctk.CTkLabel(frame, text="提取PDF文本内容，保存为可编辑的Word文档",
                     font=ctk.CTkFont(size=13), text_color="#666").pack(anchor="w", pady=(0, 10))

        file_frame = ctk.CTkFrame(frame, fg_color="transparent")
        file_frame.pack(fill="x", pady=5)
        self.word_path = ctk.StringVar()
        ctk.CTkEntry(file_frame, textvariable=self.word_path).pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(file_frame, text="选择PDF", command=self._select_word_pdf, width=100).pack(side="left")

        ctk.CTkButton(frame, text="📄 一键转Word", command=self._convert_to_word, height=40).pack(pady=15)

        self.word_log = ctk.CTkTextbox(frame, height=200)
        self.word_log.pack(fill="both", expand=True, pady=5)
        self.log(self.word_log, "✅ 转换为可编辑Word文档 (.docx)")

    def _select_word_pdf(self):
        path = filedialog.askopenfilename(filetypes=[("PDF文件", "*.pdf")])
        if path:
            self.word_path.set(path)

    def _convert_to_word(self):
        pdf_path = self.word_path.get()
        if not pdf_path:
            messagebox.showerror("错误", "请先选择PDF文件！")
            return

        save_path = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word文档", "*.docx")],
            initialfile="转换结果.docx"
        )
        if not save_path:
            return

        self.log(self.word_log, "正在转换...")
        try:
            doc = Document()
            # 优先使用 PyMuPDF 提取文本（中文支持更好）
            pdf_doc = fitz.open(pdf_path)
            for i, page in enumerate(pdf_doc):
                page_text = page.get_text()
                if page_text and page_text.strip():
                    doc.add_paragraph(f"===== 第 {i+1} 页 =====")
                    doc.add_paragraph(page_text.strip())
            pdf_doc.close()

            # 如果 PyMuPDF 没提取到内容，回退到 PyPDF2
            if len(doc.paragraphs) == 0:
                reader = PdfReader(pdf_path)
                for i, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text:
                        doc.add_paragraph(f"===== 第 {i+1} 页 =====")
                        doc.add_paragraph(text)

            doc.save(save_path)
            self.log(self.word_log, "转换成功！")
            messagebox.showinfo("成功", "PDF转Word完成！")
        except Exception as e:
            self.log(self.word_log, f"失败：{str(e)}")
            messagebox.showerror("失败", f"转换出错：{str(e)}")
